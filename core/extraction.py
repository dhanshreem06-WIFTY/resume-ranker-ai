from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from PyPDF2 import PdfReader

SUPPORTED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".txt", ".md", ".csv", ".rtf", ".odt",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = clean_text("\n".join(pages))
    if len(re.sub(r"\s+", "", text)) >= 80:
        return text

    # Scanned-PDF fallback: render pages and OCR them with RapidOCR.
    try:
        import fitz
        from core.ocr import ocr_image_bytes

        doc = fitz.open(stream=data, filetype="pdf")
        ocr_pages = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
            ocr_pages.append(ocr_image_bytes(pix.tobytes("png")))
        return clean_text("\n\n".join(ocr_pages))
    except Exception as exc:
        if not text:
            raise ValueError(f"Scanned PDF OCR could not run: {exc}") from exc
        return text


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return clean_text("\n".join(parts))


def extract_text_file(data: bytes) -> str:
    return clean_text(data.decode("utf-8", errors="ignore"))


def extract_rtf(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text
    return clean_text(rtf_to_text(data.decode("utf-8", errors="ignore")))


def extract_odt(data: bytes) -> str:
    from odf import text as odf_text
    from odf.opendocument import load

    with tempfile.NamedTemporaryFile(suffix=".odt", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        doc = load(tmp_path)
        parts = []
        for element in doc.getElementsByType(odf_text.P):
            parts.append("".join(str(node) for node in element.childNodes))
        return clean_text("\n".join(parts))
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


def extract_doc_legacy(data: bytes) -> str:
    """Extract old .doc using LibreOffice when available (also used by Docker deployment)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError("Legacy .doc files need LibreOffice installed. Use .docx, PDF, or an image if LibreOffice is unavailable.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "resume.doc"
        src.write_bytes(data)
        cmd = [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, str(src)]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        out = Path(tmp) / "resume.txt"
        if result.returncode != 0 or not out.exists():
            raise ValueError(f"Could not convert .doc file: {result.stderr.strip() or 'LibreOffice conversion failed'}")
        return clean_text(out.read_text(encoding="utf-8", errors="ignore"))


def extract_image(data: bytes) -> str:
    from core.ocr import ocr_image_bytes
    return clean_text(ocr_image_bytes(data))


def extract_resume(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or 'no extension'}")
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext == ".doc":
        return extract_doc_legacy(data)
    if ext in {".txt", ".md", ".csv"}:
        return extract_text_file(data)
    if ext == ".rtf":
        return extract_rtf(data)
    if ext == ".odt":
        return extract_odt(data)
    if ext in IMAGE_EXTENSIONS:
        return extract_image(data)
    raise ValueError(f"Unsupported file type: {ext}")
